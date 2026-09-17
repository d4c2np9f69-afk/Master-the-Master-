#!/usr/bin/env python3
"""Build bid/BID_COMPARISON.xlsx - the printable HVAC bid comparison.

Rows are Jeff's OWN spec, taken from the bid request he sent contractors on 2026-09-16.
Columns are the bidders. Regenerate with:  python3 scripts/make-bid-sheet.py

Filled cells are BLUE and come from a written quote or email - nothing here is from memory.
Empty cells are YELLOW: fill them in at the table. RED rows are deal breakers.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = "Arial"
wb = Workbook(); ws = wb.active; ws.title = "Bid Comparison"

hdr  = Font(name=F, size=11, bold=True, color="FFFFFF")
secf_font = Font(name=F, size=10, bold=True, color="FFFFFF")
bod  = Font(name=F, size=9)
red  = Font(name=F, size=9, bold=True, color="C00000")
blue = Font(name=F, size=9, color="0000FF")

navy   = PatternFill("solid", fgColor="1F3864")
secfil = PatternFill("solid", fgColor="4472C4")
yellow = PatternFill("solid", fgColor="FFFF00")
grey   = PatternFill("solid", fgColor="F2F2F2")
pink   = PatternFill("solid", fgColor="FFD9D9")

thin = Side(style="thin", color="BFBFBF")
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ctr  = Alignment(horizontal="center", vertical="center", wrap_text=True)

BIDDERS = ["Daniels\n(benchmark)", "Peters H&A\n(Tyler O'Berry)", "Goodlettsville\n(Bill)",
           "Butler A/C\n(Aaron)", "Hunter\n(Mark/Daniel)", "______________"]
N = 1 + len(BIDDERS)
r = 1

def merged(text, font, fill=None, h=None):
    global r
    c = ws.cell(r, 1, text); c.font = font; c.alignment = wrap
    if fill: c.fill = fill
    for i in range(1, N): ws.cell(r, 1 + i).fill = fill or PatternFill()
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N)
    if h: ws.row_dimensions[r].height = h
    r += 1

merged("HVAC BID COMPARISON  —  301 S Aztec Dr, White House TN 37188", Font(name=F, size=14, bold=True))
merged("2.5-ton gas/electric package unit · R-454B · 995 CFM design airflow · quotes due Thu 2026-09-17", Font(name=F, size=9, italic=True))
r += 1
merged("LEGEND:   yellow = fill in at the table    ·    blue = taken from a written quote or email    ·    RED = hard requirement, a 'no' disqualifies",
       Font(name=F, size=9, bold=True), grey)
r += 1

hrow = r
c = ws.cell(r, 1, "REQUIREMENT"); c.font = hdr; c.fill = navy; c.alignment = ctr; c.border = box
for i, b in enumerate(BIDDERS):
    c = ws.cell(r, 2 + i, b); c.font = hdr; c.fill = navy; c.alignment = ctr; c.border = box
ws.row_dimensions[r].height = 32
r += 1

def section(t):
    global r
    c = ws.cell(r, 1, t); c.font = secf_font; c.fill = secfil; c.alignment = wrap; c.border = box
    for i in range(1, N):
        cc = ws.cell(r, 1 + i); cc.fill = secfil; cc.border = box
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N); r += 1

def row(label, vals=None, hard=False, h=None):
    global r
    c = ws.cell(r, 1, label); c.font = red if hard else bod; c.alignment = wrap; c.border = box
    if hard: c.fill = pink
    vals = vals or {}
    for i in range(len(BIDDERS)):
        cell = ws.cell(r, 2 + i, vals.get(i, "")); cell.border = box; cell.alignment = wrap
        if vals.get(i): cell.font = blue
        else: cell.font = bod; cell.fill = yellow
    if h: ws.row_dimensions[r].height = h
    r += 1

OK = "YES - "

section("WHO THEY ARE")
row("Contact / title", {1: "Tyler O'Berry, Manager/Sales", 2: "Bill (office: Candy Faulkner)", 3: "Aaron Lee", 4: "Mark / Daniel"})
row("Phone", {1: "Office (615) 757-5497\nCell (629) 292-1310", 2: "(615) 479-0886"})
row("Email", {1: "TOberry@petershvac.net", 2: "bill.ghandc@gmail.com", 3: "butleracservice@yahoo.com", 4: "info@huntertn.com"})
row("TN LICENSE NUMBER", {1: "#81099 - MECHANICAL"}, hard=True)

section("THE UNIT")
row("Exact model number", {0: "Carrier 48NL-B300603", 1: "Carrier 48NL-B300603"})
row("Tonnage / gas BTU input / output", {0: "2.5 ton - 60,000 BTU", 1: "2.5 ton - 60,000 BTU gas.  INPUT vs OUTPUT still not split"})
row("SEER2", {0: "13.4", 1: "13.4"})
row("Single stage or 2-stage", {0: "single", 1: "single, ECM blower.  The 15.2 SEER2 2-stage question is STILL UNANSWERED - no price, no recommendation"}, h=40)
row("Refrigerant", {0: "R-454B"})
row("HEAT EXCHANGER MATERIAL - MUST BE STAINLESS STEEL", {0: "lifetime stainless", 1: "STAINLESS STEEL, in writing TWICE: 'heat exchanger will be stainless steel per customer request' and 'LIFETIME HEAT EXCHANGER (STAINLESS STEEL)'"}, hard=True, h=52)
row("Heat exchanger / parts / compressor warranty", {0: "lifetime HX (register in 90 days or drops to 20 yr)", 1: "LIFETIME HX + 10 yr parts + 10 yr compressor"}, h=34)
row("Labor warranty - term", {0: "10 yr", 1: "5 yr INCLUDED (was 2 yr on the first proposal)"})
row("Labor warranty - cost to reach 10 yr", {0: "$800", 1: "+$500.00"})
row("Labor warranty - what it covers", {0: "all labour + freon", 1: "all labour on warranty parts repairs, priority service, discounted maintenance"}, h=34)

section("DUCTWORK - every item below is IN the price")
row("1. Twist at unit REMOVED, new square-to-round on BOTH supply and return", {0: "included", 1: OK + "'run the new R/A trunk in a strait shot to ELIMINATE the existing cross over/twist' + square-to-round take-offs on BOTH return and supply"}, hard=True, h=52)
row("2. Dining / master / living upgraded 6\" to 8\" (saddle tap, damper, boot)", {1: OK + "all three, with dampers on every new/modified lead for balancing"}, h=34)
row("3. New 7\" run to garage, cored through brick and sealed, brown register", {1: OK + "7\" lead + a few ft of 12\" hard pipe, chased through the brick, metal critter guard inside the crawlspace", 3: "REFUSES - says a branch duct into a garage violates code (fume backfeed) and they will not do the job"}, h=52)
row("4. New 16\" return with BOTH sheet-metal 90s", {1: "PARTIAL - the (2) 16\" 90s ARE galvanized sheet metal, but the TRUNK is 16\" FLEX, not sheet metal"}, h=40)
row("4b. 18\" return priced as a SEPARATE option", {1: "+$160.00"})
row("5. Return filter grille 20x25 STAYS", {0: "n/a - stays", 1: "stays, plus 6 months of 20x25x1 pleated filters free"})
row("6. Existing duct inspected, re-strapped, sealed (NOT replaced for damp insulation)", {1: OK + "full mastic seal, mastic-backed metal tape, re-strapped high and tight, NEW sheet metal return panning sealed + insulated, return cavity sealed, microbial spray"}, h=52)
row("7. Unit base closed off at the pad - rodent / snake / weather", {1: OK + "new shroud, critter guards, sealed to the structure, PVC drain with P-trap, new plastic pad"}, h=40)
row("8. UVC REKO R2000 (Jeff supplies) installed + wired, lamps facing coil, no extra labour", {1: OK + "installed facing the coil, NO additional charge"}, h=34)
row("9. ecobee EB-STATE6P-01 (Jeff supplies): new 18 AWG + C wire, install + start-up", {1: OK + "install, set-up AND programming at no cost. 18/8 wire pull is $295"}, h=40)

section("COMMERCIAL TERMS")
row("Ductwork itemised on the quote", {1: OK + "3-page proposal with a full line-item breakdown"})
row("Price per linear foot")
row("Permit included", {1: "$95.00"})
row("Start-up readings ON THE INVOICE: static pressure + airflow set to 995 CFM", {1: "Static pressure recorded at completion.  '995 CFM' appears NOWHERE in the document"}, h=40)
row("How soon could you start")

section("PRICE")
row("Equipment + install", {0: "$8,200", 1: "$8,520.00"})
row("Ductwork", {0: "included in $9,000", 1: "$2,750.00"})
row("Options / add-ons", {0: "+$800 = 10-yr labour & freon", 1: "+$500 10-yr labour | +$160 18\" return | -$300 TVA rebate"}, h=34)
tot = r
row("TOTAL DELIVERED", {0: "$9,000 delivered", 1: "$11,270.00  (8,520 + 2,750)"})
for i in range(N):
    ws.cell(tot, 1 + i).font = Font(name=F, size=11, bold=True)

r += 1
section("READ THIS BEFORE COMPARING THE TOTALS")
for n in [
    "THE TWO TOTALS ARE NOT THE SAME JOB. Daniels' $9,000 covers the unit, the twist fix and new supply/return flex. Peters' $11,270 ALSO includes three 6\"-to-8\" branch upgrades, the 7\" garage run, a new 16\" return trunk with sheet-metal 90s, a full mastic duct seal, new sheet metal return panning, microbial treatment, the permit and six months of filters. Compare scope before price.",
    "Peters has NO printed grand total. $11,270.00 is arithmetic, and it checks two independent ways: 8,520 + 2,750 = 11,270, and the page-3 breakdown (6,260 unit w/ pad + 2,470 labour + 1,475 ducting materials + 325 duct seal + 350 misc + 95 permit + 295 t-stat wire) sums to exactly 11,270. Both sections state the total includes all sales tax.",
    "LIKE FOR LIKE ON LABOUR: Daniels $9,000 already has 10-yr labour+freon. Peters needs +$500 to reach 10 yr, so the comparable Peters figure is $11,770 before the -$300 TVA rebate.",
    "STILL OWED BY DANIELS, logged 2026-09-10 and still open: the flex diameters IN WRITING. The revised proposal deleted the 16\" flex supply/return, the 16\"-14\" reducer and the 875 CFM line. That connection work is the entire reason $9,000 counts as turnkey. Do not sign without it.",
    "STILL OWED BY PETERS: the 15.2 SEER2 two-stage price and his recommendation (asked, not answered); gas BTU input vs output; and whether the 16\" return TRUNK can be sheet metal rather than flex.",
    "MINOR BUT IT IS IN THE CONTRACT YOU WOULD SIGN: Peters' 'Notice to Owner' boilerplate cites MISSOURI lien law (Chapter 429, RSMO) in a Tennessee contract. Almost certainly a template leftover - worth one question.",
    "Peters proposal dated 9/17/2026, valid 90 days, prepared by Tyler O'Berry. Payment due 30 days after completion; 2% surcharge on card payments of $1,000 or more.",
    "Carrier lifetime stainless heat exchanger must be REGISTERED WITHIN 90 DAYS of install or it drops to 20 years.",
]:
    c = ws.cell(r, 1, "- " + n); c.font = Font(name=F, size=8); c.alignment = wrap
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N)
    ws.row_dimensions[r].height = 30; r += 1

ws.column_dimensions['A'].width = 46
for i in range(len(BIDDERS)):
    ws.column_dimensions[get_column_letter(2 + i)].width = 25
ws.freeze_panes = "B%d" % (hrow + 1)
ws.print_title_rows = "%d:%d" % (hrow, hrow)
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.25
ws.page_margins.top = ws.page_margins.bottom = 0.35

out = "bid/BID_COMPARISON.xlsx"
wb.save(out)
print("WROTE", out)
